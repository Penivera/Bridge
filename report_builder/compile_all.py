import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

import report_builder.prelims as prelims
import report_builder.ch1 as ch1
import report_builder.ch2 as ch2
import report_builder.ch3 as ch3
import report_builder.ch4 as ch4

def build_markdown():
    print("Generating SIWES_TECHNICAL_REPORT_AKIRS_BRIDGE.md...")
    full_md = (
        prelims.TITLE_PAGE_MD + "\n\n" +
        ch1.CHAPTER_ONE_MD + "\n\n" +
        ch2.CHAPTER_TWO_MD + "\n\n" +
        ch3.CHAPTER_THREE_MD + "\n\n" +
        ch4.CHAPTER_FOUR_MD
    )
    with open("SIWES_TECHNICAL_REPORT_AKIRS_BRIDGE.md", "w", encoding="utf-8") as f:
        f.write(full_md)
    print(f"Markdown report generated successfully! Total characters: {len(full_md)}")
    return full_md

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_styled_table(doc, headers, rows):
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    
    # Format header
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h.strip()
        set_cell_background(hdr_cells[i], "1F4E79") # Deep Blue
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=180, right=180)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.bold = True
            run.font.name = "Calibri"
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(255, 255, 255)
            
    # Format data rows
    for r_idx, row_data in enumerate(rows):
        row_cells = table.rows[r_idx + 1].cells
        bg_color = "F2F5F8" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].text = val.strip()
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=100, bottom=100, left=150, right=150)
            p = row_cells[c_idx].paragraphs[0]
            if c_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.name = "Calibri"
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(40, 40, 40)

    # Empty paragraph after table
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(4)
    p_after.paragraph_format.space_after = Pt(8)

def add_code_or_ascii_block(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.2)
    cell = table.rows[0].cells[0]
    set_cell_background(cell, "F4F5F7")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Border
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="0" w:color="D0D5DD"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="D0D5DD"/><w:left w:val="single" w:sz="18" w:space="0" w:color="1F4E79"/><w:right w:val="single" w:sz="4" w:space="0" w:color="D0D5DD"/></w:tcBorders>')
    tcPr.append(borders)
    
    cell.text = ""
    lines = text.strip("\n").split("\n")
    for idx, l in enumerate(lines):
        p = cell.paragraphs[0] if idx == 0 else cell.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.05
        run = p.add_run(l)
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(30, 41, 59)
        
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_after = Pt(6)

def build_docx(markdown_content):
    print("Compiling SIWES_TECHNICAL_REPORT_AKIRS_BRIDGE.docx...")
    doc = Document()
    
    # Page setup
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.25) # 1.25 inch for academic binding
        s.right_margin = Inches(1.0)
        
    lines = markdown_content.split("\n")
    i = 0
    in_code_block = False
    code_buffer = []
    
    while i < len(lines):
        line = lines[i]
        
        # Check code block fences
        if line.strip().startswith("```"):
            if in_code_block:
                add_code_or_ascii_block(doc, "\n".join(code_buffer))
                code_buffer = []
                in_code_block = False
            else:
                in_code_block = True
                code_buffer = []
            i += 1
            continue
            
        if in_code_block:
            code_buffer.append(line)
            i += 1
            continue
            
        # Check Markdown tables
        if line.strip().startswith("|") and line.strip().endswith("|"):
            # Table detected
            table_lines = [line.strip()]
            i += 1
            while i < len(lines) and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
                
            if len(table_lines) >= 3 and "---" in table_lines[1]:
                headers = [c.strip() for c in table_lines[0].split("|")[1:-1]]
                rows = []
                for row_str in table_lines[2:]:
                    cols = [c.strip() for c in row_str.split("|")[1:-1]]
                    rows.append(cols)
                add_styled_table(doc, headers, rows)
            continue
            
        # Page breaks
        if "<div style=\"page-break-after: always;\"></div>" in line:
            doc.add_page_break()
            i += 1
            continue
            
        # Separator line
        if line.strip() == "---":
            i += 1
            continue

        # Headings
        if line.startswith("# "):
            text = line[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(text)
            run.font.name = "Calibri"
            run.font.size = Pt(18)
            run.font.bold = True
            run.font.color.rgb = RGBColor(15, 23, 42) # Dark Slate
            i += 1
            continue
            
        if line.startswith("## "):
            text = line[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(text)
            run.font.name = "Calibri"
            run.font.size = Pt(15)
            run.font.bold = True
            run.font.color.rgb = RGBColor(31, 78, 121) # Primary Blue
            i += 1
            continue

        if line.startswith("### "):
            text = line[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(text)
            run.font.name = "Calibri"
            run.font.size = Pt(12.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(45, 55, 72)
            i += 1
            continue

        if line.startswith("#### "):
            text = line[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(text)
            run.font.name = "Calibri"
            run.font.size = Pt(11.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(71, 85, 105)
            i += 1
            continue
            
        if line.startswith("##### "):
            text = line[6:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(text)
            run.font.name = "Calibri"
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.italic = True
            run.font.color.rgb = RGBColor(71, 85, 105)
            i += 1
            continue

        # Bullet points
        if line.strip().startswith("* ") or line.strip().startswith("- "):
            text = line.strip()[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            # Parse bold runs inside bullets
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    r = p.add_run(part)
                r.font.name = "Calibri"
                r.font.size = Pt(11)
                r.font.color.rgb = RGBColor(51, 51, 51)
            i += 1
            continue
            
        # Numbered list
        m_num = re.match(r'^(\d+)\.\s+(.*)', line.strip())
        if m_num:
            text = m_num.group(2).strip()
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                else:
                    r = p.add_run(part)
                r.font.name = "Calibri"
                r.font.size = Pt(11)
                r.font.color.rgb = RGBColor(51, 51, 51)
            i += 1
            continue

        # Standard paragraph
        if line.strip():
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.line_spacing = 1.18
            parts = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', line)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.font.bold = True
                elif part.startswith("*") and part.endswith("*"):
                    r = p.add_run(part[1:-1])
                    r.font.italic = True
                elif part.startswith("`") and part.endswith("`"):
                    r = p.add_run(part[1:-1])
                    r.font.name = "Consolas"
                    r.font.size = Pt(9.5)
                    r.font.color.rgb = RGBColor(180, 40, 40)
                else:
                    r = p.add_run(part)
                r.font.name = "Calibri"
                r.font.size = Pt(11)
                r.font.color.rgb = RGBColor(40, 40, 40)
        i += 1

    docx_path = "SIWES_TECHNICAL_REPORT_AKIRS_BRIDGE.docx"
    doc.save(docx_path)
    print(f"Word document compiled successfully: {docx_path}")

if __name__ == "__main__":
    md_content = build_markdown()
    build_docx(md_content)
