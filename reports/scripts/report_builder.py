#!/usr/bin/env python3
"""
Master Builder Script: Professional Formatting of SIWES Technical Report for Ben, Peniel Sunday.
Strictly references target-doc.pdf for structure, typography, and styling.
"""

import sys
import os
import re
from pathlib import Path

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

WORKSPACE = Path("/home/peni/Projects/bridge")
EXTRACTED_TXT = WORKSPACE / "extracted_siwes_report.txt"
WORKING_DOCX = WORKSPACE / "SIWES_Technical_Report_Peniel.docx"
FINAL_DOCX = WORKSPACE / "SIWES_Technical_Report_Peniel_Final.docx"
FINAL_PDF = WORKSPACE / "SIWES_Technical_Report_Peniel_Final.pdf"

# Palette colors from design system
NAVY_HEX = "0B2545"
TEAL_HEX = "13678A"
SLATE_HEX = "45596C"
BORDER_HEX = "C5CED4"
BG_LIGHT_HEX = "D9E8EC"
WHITE_HEX = "FFFFFF"


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set inner cell padding in twips."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        r'<w:tcMar %s>'
        r'  <w:top w:w="%d" w:type="dxa"/>'
        r'  <w:bottom w:w="%d" w:type="dxa"/>'
        r'  <w:left w:w="%d" w:type="dxa"/>'
        r'  <w:right w:w="%d" w:type="dxa"/>'
        r'</w:tcMar>' % (nsdecls('w'), top, bottom, left, right)
    )
    tcPr.append(tcMar)


def set_cell_shading(cell, color_hex):
    """Set background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(r'<w:shd %s w:fill="%s"/>' % (nsdecls('w'), color_hex))
    tcPr.append(shd)


def add_table_borders(table):
    """Apply elegant academic borders to a table."""
    borders = parse_xml(
        r'<w:tblBorders %s>'
        r'  <w:top w:val="single" w:sz="6" w:space="0" w:color="0B2545"/>'
        r'  <w:left w:val="none"/>'
        r'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="0B2545"/>'
        r'  <w:right w:val="none"/>'
        r'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="C5CED4"/>'
        r'  <w:insideV w:val="none"/>'
        r'</w:tblBorders>' % nsdecls('w')
    )
    table._tbl.tblPr.append(borders)


def format_table_header(row, col_names):
    """Format the header row of a table."""
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(r'<w:tblHeader %s/>' % nsdecls('w')))
    trPr.append(parse_xml(r'<w:cantSplit %s/>' % nsdecls('w')))
    for i, name in enumerate(col_names):
        cell = row.cells[i]
        set_cell_shading(cell, BG_LIGHT_HEX)
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(name)
        run.font.name = "Times New Roman"
        run.font.size = Pt(10)
        run.font.bold = True
        run.font.color.rgb = RGBColor(11, 37, 69)


def add_cover_page(doc):
    """Construct a perfectly balanced cover page matching target-doc.pdf strictly."""
    s = doc.sections[0]
    s.page_width = Inches(8.27)
    s.page_height = Inches(11.69)
    s.top_margin = Inches(1.0)
    s.bottom_margin = Inches(1.0)
    s.left_margin = Inches(1.0)
    s.right_margin = Inches(1.0)

    def cp_para(text, size=12, bold=True, before=0, after=0):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(after)
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = RGBColor(0, 0, 0)
        return p

    cp_para("REPORT ON STUDENT INDUSTRIAL WORK EXPERIENCE SCHEME (SIWES)", 12, True, 0, 3)
    cp_para("EXPERIENCES AND TECHNICAL PROJECT UNDERTAKEN AT", 11, True, 0, 3)
    cp_para("AKWA IBOM STATE INTERNAL REVENUE SERVICE (AKIRS)", 12, True, 0, 3)
    cp_para("(JOINT REVENUE BOARD DIGITAL INFRASTRUCTURE PROGRAMME)", 11, True, 0, 3)
    cp_para("REVENUE HOUSE, UYO, AKWA IBOM STATE", 11, True, 0, 24)

    cp_para("ON", 11, True, 0, 6)
    cp_para("BRIDGE: A FAULT-TOLERANT DISTRIBUTED INGRESS ROUTING SYSTEM AND REVENUE PLATFORM INFRASTRUCTURE", 13, True, 0, 28)

    cp_para("BY", 11, True, 0, 6)
    cp_para("BEN, PENIEL SUNDAY", 13, True, 0, 3)
    cp_para("REG NO: 22/SC/CO/1181", 12, True, 0, 3)
    cp_para("DEPARTMENT OF COMPUTER SCIENCE", 11, True, 0, 3)
    cp_para("FACULTY OF COMPUTING", 11, True, 0, 3)
    cp_para("UNIVERSITY OF UYO, UYO, AKWA IBOM STATE", 11, True, 0, 24)

    cp_para("COURSE CODE: CSC 329", 12, True, 0, 24)

    cp_para("SUBMITTED TO:", 11, True, 0, 6)
    cp_para("DEPARTMENT OF COMPUTER SCIENCE", 11, True, 0, 3)
    cp_para("FACULTY OF COMPUTING", 11, True, 0, 3)
    cp_para("UNIVERSITY OF UYO, UYO, AKWA IBOM STATE", 11, True, 0, 28)

    cp_para("IN PARTIAL FULFILMENT OF THE REQUIREMENT FOR THE AWARD", 11, True, 0, 3)
    cp_para("OF A BACHELOR OF SCIENCE (B.SC) DEGREE IN COMPUTER SCIENCE", 11, True, 0, 32)

    cp_para("OCTOBER, 2026", 12, True, 0, 0)


def setup_toc_section(doc):
    """Add Section 2 for Table of Contents with Roman numeral pagination."""
    s2 = doc.add_section(WD_SECTION.NEW_PAGE)
    s2.page_width = Inches(8.27)
    s2.page_height = Inches(11.69)
    s2.top_margin = Inches(1.0)
    s2.bottom_margin = Inches(1.0)
    s2.left_margin = Inches(1.0)
    s2.right_margin = Inches(1.0)
    s2.header.is_linked_to_previous = False
    s2.footer.is_linked_to_previous = False

    pg_num_type = parse_xml(r'<w:pgNumType %s w:fmt="lowerRoman" w:start="1"/>' % nsdecls('w'))
    s2._sectPr.append(pg_num_type)

    fp = s2.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.line_spacing = 1.0
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)
    fld = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
    fp._p.append(fld)
    return s2


def add_toc_content(doc, toc_data):
    """Populate Table of Contents with exact tab leader dots and accurate page numbers."""
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(12)
    p_title.paragraph_format.space_after = Pt(18)
    r_title = p_title.add_run("Table of Content")
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(14)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0, 0, 0)

    # Insert Word TOC field
    p_fld = doc.add_paragraph()
    p_fld.paragraph_format.space_before = Pt(0)
    p_fld.paragraph_format.space_after = Pt(0)
    fld_toc = parse_xml(r'<w:fldSimple %s w:instr="TOC \o &quot;1-3&quot; \h \z \u"/>' % nsdecls('w'))
    p_fld._p.append(fld_toc)

    # Pre-populated entries matching target-doc.pdf
    for item in toc_data:
        level, title_text, page_num = item
        p = doc.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(Inches(6.27), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        p.paragraph_format.line_spacing = 1.15
        
        # Formatting per level
        if level == "chapter":
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.left_indent = Inches(0.0)
            r1 = p.add_run(title_text)
            r1.font.name = "Times New Roman"
            r1.font.size = Pt(11)
            r1.font.bold = True
            r2 = p.add_run(f"\t{page_num}")
            r2.font.name = "Times New Roman"
            r2.font.size = Pt(11)
            r2.font.bold = True
        elif level == "h1":
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.left_indent = Inches(0.0)
            r1 = p.add_run(title_text)
            r1.font.name = "Times New Roman"
            r1.font.size = Pt(10.5)
            r1.font.bold = True
            r2 = p.add_run(f"\t{page_num}")
            r2.font.name = "Times New Roman"
            r2.font.size = Pt(10.5)
            r2.font.bold = True
        elif level == "h2":
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.left_indent = Inches(0.15)
            r1 = p.add_run(title_text)
            r1.font.name = "Times New Roman"
            r1.font.size = Pt(10)
            r1.font.bold = False
            r2 = p.add_run(f"\t{page_num}")
            r2.font.name = "Times New Roman"
            r2.font.size = Pt(10)
            r2.font.bold = False
        elif level == "h3":
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.left_indent = Inches(0.35)
            r1 = p.add_run(title_text)
            r1.font.name = "Times New Roman"
            r1.font.size = Pt(9.5)
            r1.font.bold = False
            r2 = p.add_run(f"\t{page_num}")
            r2.font.name = "Times New Roman"
            r2.font.size = Pt(9.5)
            r2.font.bold = False


def setup_body_section(doc):
    """Add Section 3 for Chapters with Arabic numeral pagination starting at 1."""
    s3 = doc.add_section(WD_SECTION.NEW_PAGE)
    s3.page_width = Inches(8.27)
    s3.page_height = Inches(11.69)
    s3.top_margin = Inches(1.0)
    s3.bottom_margin = Inches(1.0)
    s3.left_margin = Inches(1.0)
    s3.right_margin = Inches(1.0)
    s3.header.is_linked_to_previous = False
    s3.footer.is_linked_to_previous = False

    pg_num_type = parse_xml(r'<w:pgNumType %s w:fmt="decimal" w:start="1"/>' % nsdecls('w'))
    s3._sectPr.append(pg_num_type)

    fp = s3.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.line_spacing = 1.0
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)
    fld = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
    fp._p.append(fld)
    return s3


def add_chapter_header(doc, chapter_num_str, chapter_title_str):
    """Add academic chapter heading following target-doc.pdf strictly."""
    p_num = doc.add_paragraph()
    p_num.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_num.paragraph_format.space_before = Pt(12)
    p_num.paragraph_format.space_after = Pt(2)
    p_num.paragraph_format.keep_with_next = True
    r_num = p_num.add_run(chapter_num_str)
    r_num.font.name = "Times New Roman"
    r_num.font.size = Pt(14)
    r_num.font.bold = True
    r_num.font.color.rgb = RGBColor(0, 0, 0)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(2)
    p_title.paragraph_format.space_after = Pt(18)
    p_title.paragraph_format.keep_with_next = True
    r_title = p_title.add_run(chapter_title_str)
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(14)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0, 0, 0)


def add_heading_1(doc, text):
    """Add Level 1 heading (e.g. 1.0 INTRODUCTION) 12pt Bold Left."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p


def add_heading_2(doc, text):
    """Add Level 2 heading (e.g. 1.2 BRIEF HISTORY...) 12pt Bold Left."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p


def add_heading_3(doc, text):
    """Add Level 3 heading (e.g. 2.3.1 Directorate of ICT...) 12pt Bold Left."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p


def add_body_paragraph(doc, text, bold_prefix=None):
    """Add body paragraph: Times New Roman, 12 pt, double line spacing, justified."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 2.0
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(12)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(0, 0, 0)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p


def add_bullet_point(doc, text, bold_prefix=None):
    """Add bullet item: Times New Roman, 12 pt, 1.5 line spacing, clean indent."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.first_line_indent = Inches(-0.2)
    r_sym = p.add_run("•  ")
    r_sym.font.name = "Times New Roman"
    r_sym.font.size = Pt(12)
    r_sym.font.bold = True
    r_sym.font.color.rgb = RGBColor(11, 37, 69)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(12)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(0, 0, 0)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p


def add_figure_image(doc, img_path, caption_text, width_in=5.8):
    """Add figure centered with academic caption below figure."""
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(10)
    p_img.paragraph_format.space_after = Pt(4)
    p_img.paragraph_format.keep_with_next = True
    r_img = p_img.add_run()
    r_img.add_picture(str(img_path), width=Inches(width_in))

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(4)
    p_cap.paragraph_format.space_after = Pt(14)
    p_cap.paragraph_format.keep_with_next = False
    
    parts = caption_text.split(":", 1)
    if len(parts) == 2:
        r_lbl = p_cap.add_run(parts[0] + ":")
        r_lbl.font.name = "Times New Roman"
        r_lbl.font.size = Pt(11)
        r_lbl.font.bold = True
        r_lbl.font.color.rgb = RGBColor(11, 37, 69)

        r_txt = p_cap.add_run(parts[1])
        r_txt.font.name = "Times New Roman"
        r_txt.font.size = Pt(11)
        r_txt.font.italic = True
        r_txt.font.color.rgb = RGBColor(0, 0, 0)
    else:
        r_txt = p_cap.add_run(caption_text)
        r_txt.font.name = "Times New Roman"
        r_txt.font.size = Pt(11)
        r_txt.font.bold = True


def add_code_block(doc, code_text):
    """Add formatted code block in Consolas 9.5 pt with shaded background."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.right_indent = Inches(0.3)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.15
    
    # Border & shading via XML
    pBdr = parse_xml(
        r'<w:pBdr %s>'
        r'  <w:left w:val="single" w:sz="18" w:space="10" w:color="0B2545"/>'
        r'</w:pBdr>' % nsdecls('w')
    )
    shd = parse_xml(r'<w:shd %s w:fill="EEF1F4"/>' % nsdecls('w'))
    p._p.get_or_add_pPr().append(pBdr)
    p._p.get_or_add_pPr().append(shd)

    r = p.add_run(code_text)
    r.font.name = "Consolas"
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(26, 31, 36)


def add_reference_entry(doc, ref_text):
    """Add APA 7th style reference entry with 0.5-inch hanging indent."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 2.0
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.5)
    
    r = p.add_run(ref_text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0, 0, 0)


print("Master Builder core framework defined.")
